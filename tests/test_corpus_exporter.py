from pathlib import Path

import pytest
from openpyxl import load_workbook

from complaint_dedup.async_database import AsyncDatabase
from complaint_dedup.corpus_exporter import export_corpus
from complaint_dedup.corpus_pipeline import CorpusProcessor
from complaint_dedup.corpus_repository import CorpusRepository
from complaint_dedup.pipeline import InputRecord


def record(row: int, address: str, title: str) -> InputRecord:
    return InputRecord(
        source="B",
        source_row=row,
        work_order_id=f"WO-{row}",
        received_at=f"2026-08-12 0{row}:00:00",
        title=title,
        category="道路积水",
        category_level_4="道路积水",
        appeal_text=f"地址：{address}。\n事项：道路积水。",
        raw_fields={"工单编号": f"WO-{row}", "诉求标题": title, "回复内容": "=1+1"},
    )


def unknown_location_record(row: int) -> InputRecord:
    return InputRecord(
        source="B",
        source_row=row,
        work_order_id=f"WO-{row}",
        received_at=f"2026-08-12 0{row}:00:00",
        title="咨询失业保险",
        category="失业保险咨询",
        category_level_4="失业保险咨询",
        appeal_text="事项：咨询失业保险待遇。",
        raw_fields={"工单编号": f"WO-{row}", "诉求标题": "咨询失业保险"},
    )


@pytest.mark.asyncio
async def test_export_uses_two_sheets_original_columns_and_alternating_event_colors(
    tmp_path: Path,
):
    database = AsyncDatabase(f"sqlite+aiosqlite:///{tmp_path / 'export.db'}")
    await database.initialize()
    repository = CorpusRepository(database)
    processor = CorpusProcessor(repository)
    staged = await processor.stage_records(
        name="历史",
        batch_type="bootstrap_history",
        file_name="history.xlsx",
        file_hash="9" * 64,
        records=[
            record(2, "江海区礼乐街道德昌电机门口", "德昌积水一"),
            record(3, "江海区礼乐街道德昌电机门口", "德昌积水二"),
            record(4, "江海区礼乐街道文华豪庭北门", "文华积水一"),
            record(5, "江海区礼乐街道文华豪庭北门", "文华积水二"),
            record(6, "江海区礼乐街道中心公园南门", "公园积水"),
        ],
    )
    await processor.approve_bootstrap(
        staged.batch_id, staged.dictionary_version_id, approved_by="tester"
    )
    output = await export_corpus(repository, tmp_path / "result.xlsx")
    await database.close()

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["重复项", "孤立工单"]
    duplicate = workbook["重复项"]
    singleton = workbook["孤立工单"]
    assert [cell.value for cell in duplicate[1]][:4] == [
        "事件名称",
        "工单编号",
        "诉求标题",
        "回复内容",
    ]
    assert duplicate.freeze_panes == "A2"
    assert duplicate.auto_filter.ref == duplicate.dimensions
    assert singleton.max_row == 2
    assert duplicate["D2"].value == "'=1+1"
    assert duplicate["A2"].fill.fgColor.rgb != duplicate["A4"].fill.fgColor.rgb


@pytest.mark.asyncio
async def test_export_includes_bootstrap_records_without_an_exact_event_key(
    tmp_path: Path,
):
    database = AsyncDatabase(f"sqlite+aiosqlite:///{tmp_path / 'complete.db'}")
    await database.initialize()
    repository = CorpusRepository(database)
    processor = CorpusProcessor(repository)
    staged = await processor.stage_records(
        name="历史",
        batch_type="bootstrap_history",
        file_name="history.xlsx",
        file_hash="8" * 64,
        records=[
            record(2, "江海区礼乐街道德昌电机门口", "德昌积水"),
            unknown_location_record(3),
        ],
    )
    await processor.approve_bootstrap(
        staged.batch_id, staged.dictionary_version_id, approved_by="tester"
    )

    output = await export_corpus(repository, tmp_path / "complete.xlsx")
    await database.close()

    workbook = load_workbook(output, read_only=True)
    exported_rows = sum(
        workbook[sheet_name].max_row - 1 for sheet_name in workbook.sheetnames
    )
    assert exported_rows == 2
