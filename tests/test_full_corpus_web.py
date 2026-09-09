from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl import load_workbook

from complaint_dedup.async_database import AsyncDatabase
from complaint_dedup.config import Settings
from complaint_dedup.full_corpus import FullCorpusService
from complaint_dedup.full_corpus_web import create_full_corpus_app
from complaint_dedup.full_corpus_web import _local_date
from zoneinfo import ZoneInfo
from datetime import datetime, timezone
from threading import Event, Thread
import time


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["工单编号", "受理时间", "办结时间", "诉求标题", "事发地点", "事项分类", "市民诉求", "处理部门"])
    sheet.append(["WO-1", "2026-09-01 08:00:00", "2026-09-02 08:00:00", "德昌电机门口积水", "江海区礼乐街道德昌电机门口", "道路积水", "地址：江海区礼乐街道德昌电机门口。", "礼乐街道办事处"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_default_window_date_uses_configured_local_timezone():
    value = datetime(2026, 8, 31, 16, 30, tzinfo=timezone.utc)
    assert _local_date(value, ZoneInfo("Asia/Shanghai")).isoformat() == "2026-09-01"


def test_full_corpus_web_upload_and_create_comparison(tmp_path: Path):
    settings = Settings(database_mode="sqlite", database_path=tmp_path / "app.db", _env_file=None)
    database = AsyncDatabase(f"sqlite+aiosqlite:///{settings.database_path}")
    app = create_full_corpus_app(settings, database=database)
    with TestClient(app) as client:
        response = client.post(
            "/sync",
            files={"file": ("all.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            follow_redirects=False,
        )
        assert response.status_code == 303
        sync_job_id = response.headers["location"].split("job_id=", 1)[1]
        for _ in range(50):
            sync_job = client.get(f"/jobs/{sync_job_id}").json()
            if sync_job["status"] == "completed":
                break
            time.sleep(0.02)
        assert sync_job["status"] == "completed"
        home = client.get("/")
        assert home.status_code == 200
        assert "当前工单数" in home.text
        result = client.post(
            "/comparisons",
            data={"time_field": "completed_at", "target_from": "2026-09-02", "target_to": "2026-09-02"},
            follow_redirects=False,
        )
        assert result.status_code == 303
        comparison_job_id = result.headers["location"].split("job_id=", 1)[1]
        for _ in range(50):
            comparison_job = client.get(f"/jobs/{comparison_job_id}").json()
            if comparison_job["status"] == "completed":
                break
            time.sleep(0.02)
        assert comparison_job["status"] == "completed"
        comparison_id = comparison_job["result_json"]["comparison_id"]
        detail = client.get(f"/comparisons/{comparison_id}", follow_redirects=False)
        assert detail.status_code == 303
        detail = client.get(detail.headers["location"])
        assert detail.status_code == 200
        assert "办结时间" in detail.text
        assert "处理部门" in detail.text
        assert "礼乐街道办事处" in detail.text
        exported = client.get(f"/comparisons/{comparison_id}/export")
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.content))
        assert workbook.sheetnames == ["重复项", "孤立工单"]
        assert workbook["孤立工单"]["A1"].value == "数据侧"


def test_full_corpus_web_removed_cannot_link_route(tmp_path: Path):
    settings = Settings(database_mode="sqlite", database_path=tmp_path / "app.db", _env_file=None)
    database = AsyncDatabase(f"sqlite+aiosqlite:///{settings.database_path}")
    app = create_full_corpus_app(settings, database=database)
    with TestClient(app) as client:
        client.post(
            "/sync",
            files={"file": ("all.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            follow_redirects=False,
        )
        saved = client.post("/comparisons/demo/cannot-links", follow_redirects=False)
        assert saved.status_code == 404


def test_full_corpus_web_exposes_comparison_library(tmp_path: Path):
    settings = Settings(database_mode="sqlite", database_path=tmp_path / "app.db", _env_file=None)
    database = AsyncDatabase(f"sqlite+aiosqlite:///{settings.database_path}")
    app = create_full_corpus_app(settings, database=database)
    with TestClient(app) as client:
        client.post(
            "/sync",
            files={"file": ("all.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            follow_redirects=False,
        )
        library = client.get("/comparisons")
        assert library.status_code == 200
        assert "历史比对任务" in library.text


def test_full_corpus_job_status_remains_responsive_while_job_runs(
    tmp_path: Path, monkeypatch
):
    settings = Settings(database_mode="sqlite", database_path=tmp_path / "app.db", _env_file=None)
    database = AsyncDatabase(f"sqlite+aiosqlite:///{settings.database_path}")
    app = create_full_corpus_app(settings, database=database)
    started = Event()
    release = Event()
    original_sync = FullCorpusService.sync_records

    async def blocked_sync(self, records, *, file_name, file_hash=None):
        started.set()
        if not release.wait(timeout=5):
            raise RuntimeError("测试任务未被释放")
        return await original_sync(
            self,
            records,
            file_name=file_name,
            file_hash=file_hash,
        )

    monkeypatch.setattr(FullCorpusService, "sync_records", blocked_sync)

    with TestClient(app) as client:
        response = client.post(
            "/sync",
            files={
                "file": (
                    "all.xlsx",
                    workbook_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        job_id = response.headers["location"].split("job_id=", 1)[1]
        assert started.wait(timeout=2)

        result: dict[str, object] = {}
        finished = Event()

        def read_status() -> None:
            result["response"] = client.get(f"/jobs/{job_id}")
            finished.set()

        request_thread = Thread(target=read_status)
        request_thread.start()
        try:
            assert finished.wait(timeout=1), "后台任务阻塞了 API 任务状态查询"
            assert result["response"].status_code == 200
        finally:
            release.set()
            request_thread.join(timeout=2)
