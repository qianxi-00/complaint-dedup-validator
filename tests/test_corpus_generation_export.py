from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from complaint_dedup.async_database import AsyncDatabase
from complaint_dedup.corpus_exporter import export_corpus
from complaint_dedup.corpus_pipeline import CorpusProcessor
from complaint_dedup.corpus_repository import CorpusRepository
from complaint_dedup.corpus_models import InputRecord


def _record(source: str, row: int, received: str) -> InputRecord:
    return InputRecord(
        source=source,
        source_row=row,
        work_order_id=f"WO-{source}-{row}",
        received_at=received,
        title="德昌电机门口积水",
        category="道路积水",
        category_level_4="道路积水",
        appeal_text="地址：江海区礼乐街道德昌电机门口。\n事项：道路积水。",
        raw_fields={
            "工单编号 ": f"WO-{source}-{row}",
            "受理时间 ": received,
            "诉求标题": "德昌电机门口积水",
        },
    )


async def _bootstrap(
    processor: CorpusProcessor, *, file_hash: str, received: str
) -> str:
    staged = await processor.stage_records(
        name="历史冷启动",
        batch_type="bootstrap_history",
        file_name="history.xlsx",
        file_hash=file_hash,
        records=[_record("B", 2, received)],
    )
    await processor.approve_bootstrap(
        staged.batch_id, staged.dictionary_version_id, approved_by="test"
    )
    return staged.batch_id


@pytest.mark.asyncio
async def test_rebootstrap_same_file_hash_replaces_active_generation_and_export_scope(
    tmp_path: Path,
):
    database = AsyncDatabase(f"sqlite+aiosqlite:///{tmp_path / 'generation.db'}")
    await database.initialize()
    repository = CorpusRepository(database)
    processor = CorpusProcessor(repository)

    await _bootstrap(processor, file_hash="a" * 64, received="2026-08-12 08:00:00")
    first_event = (await repository.list_events())[0]["id"]
    await _bootstrap(processor, file_hash="a" * 64, received="2026-08-12 08:00:00")
    active_events = await repository.list_events()

    assert len(active_events) == 1
    assert active_events[0]["id"] != first_event
    active_generation = await repository.active_generation()
    assert active_generation is not None
    async with database.engine.connect() as connection:
        assert await connection.scalar(text("select count(*) from corpus_records")) == 2
        assert await connection.scalar(
            text("select count(*) from corpus_sources where generation_id=:id"),
            {"id": active_generation["id"]},
        ) == 1
        statuses = await connection.execute(
            text("select status from corpus_generations order by id")
        )
        assert [row[0] for row in statuses] == ["archived", "active"]

    output = await export_corpus(repository, tmp_path / "result.xlsx")
    workbook = load_workbook(output, read_only=False, data_only=False)
    assert workbook.sheetnames == ["重复项", "孤立工单"]
    sheet = workbook["孤立工单"]
    assert sheet.max_row == 2
    assert [sheet.cell(1, index).value for index in range(1, 6)] == [
        "数据来源",
        "事件名称",
        "工单编号 ",
        "受理时间 ",
        "诉求标题",
    ]
    assert sheet["A2"].value == "历史表"
    await database.close()
