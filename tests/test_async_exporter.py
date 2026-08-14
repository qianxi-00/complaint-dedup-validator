from pathlib import Path

import openpyxl
import pytest

from complaint_dedup.async_database import AsyncDatabase
from complaint_dedup.async_exporter import export_job


@pytest.mark.asyncio
async def test_export_contains_event_name_for_every_record(tmp_path: Path) -> None:
    database = AsyncDatabase(f"sqlite+aiosqlite:///{tmp_path / 'app.db'}")
    await database.initialize()
    await database.enqueue_job("job-1", "导出任务", mode="single", total_records=2, pipeline_version="pair_v1")
    await database.add_records(
        "job-1",
        [
            {
                "source": "S",
                "source_row": 2,
                "work_order_id": "A001",
                "received_at": "2026-01-01",
                "title": "甲公司拖欠工资",
                "region": "江海区",
            },
            {
                "source": "S",
                "source_row": 3,
                "work_order_id": "A002",
                "received_at": "2026-01-03",
                "title": "乙公司噪音",
                "region": "江海区",
            },
        ],
    )
    await database.rebuild_event_groups("job-1")

    output = await export_job(database, "job-1", tmp_path / "result.xlsx")
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)

    assert workbook.sheetnames == [
        "结果总览",
        "地区统计",
        "工单明细",
        "候选对",
        "事件组",
        "抽取失败",
        "模型失败",
    ]
    rows = list(workbook["工单明细"].iter_rows(values_only=True))
    headers = list(rows[0])
    event_name_index = headers.index("事件名称")
    assert len(rows) == 3
    assert all(row[event_name_index] for row in rows[1:])
    await database.close()


@pytest.mark.asyncio
async def test_export_preserves_raw_fields_and_distinguishes_singletons(tmp_path: Path) -> None:
    database = AsyncDatabase(f"sqlite+aiosqlite:///{tmp_path / 'app.db'}")
    await database.initialize()
    await database.enqueue_job("job-1", "导出任务", mode="single", total_records=3, pipeline_version="pair_v1")
    first, second, singleton = await database.add_records(
        "job-1",
        [
            {"source": "S", "source_row": 2, "title": "甲", "region": "江海区", "raw_json": {"自定义列": "原值甲"}},
            {"source": "S", "source_row": 3, "title": "乙", "region": "江海区", "raw_json": {"自定义列": "原值乙"}},
            {"source": "S", "source_row": 4, "title": "丙", "region": "蓬江区", "raw_json": {"自定义列": "原值丙"}},
        ],
    )
    await database.upsert_candidate_pairs(
        "job-1",
        [{"record_a_id": first, "record_b_id": second, "recall_reason": "vector"}],
    )
    pair = (await database.list_candidate_pairs("job-1"))[0]
    await database.review_pair("job-1", pair["id"], "duplicate")

    output = await export_job(database, "job-1", tmp_path / "result.xlsx")
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    rows = list(workbook["工单明细"].iter_rows(values_only=True))
    headers = list(rows[0])
    values = [dict(zip(headers, row)) for row in rows[1:]]

    assert {row["原始字段｜自定义列"] for row in values} == {"原值甲", "原值乙", "原值丙"}
    assert {row["合并状态"] for row in values if row["诉求标题"] in {"甲", "乙"}} == {"已合并"}
    assert next(row for row in values if row["诉求标题"] == "丙")["合并状态"] == "单例"
    assert next(row for row in values if row["诉求标题"] == "丙")["事件组ID"] is not None
    await database.close()


@pytest.mark.asyncio
async def test_export_region_stats_count_cross_region_pair_for_each_actual_region(tmp_path: Path) -> None:
    database = AsyncDatabase(f"sqlite+aiosqlite:///{tmp_path / 'app.db'}")
    await database.initialize()
    await database.enqueue_job("job-1", "跨区统计", mode="single", total_records=2, pipeline_version="pair_v1")
    left, right = await database.add_records(
        "job-1",
        [
            {"source": "S", "source_row": 2, "title": "甲", "region": "江海区"},
            {"source": "S", "source_row": 3, "title": "乙", "region": "蓬江区"},
        ],
    )
    await database.upsert_candidate_pairs(
        "job-1",
        [{"record_a_id": left, "record_b_id": right, "recall_reason": "vector"}],
    )
    pair = (await database.list_candidate_pairs("job-1"))[0]
    await database.save_judgements(
        "job-1", {pair["id"]: {"decision": "review", "confidence": 0.7}}
    )
    await database.rebuild_event_groups("job-1")

    output = await export_job(database, "job-1", tmp_path / "result.xlsx")
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    rows = list(workbook["地区统计"].iter_rows(values_only=True))
    stats = {row[0]: row for row in rows[1:]}

    assert stats["江海区"][4] == 1
    assert stats["蓬江区"][4] == 1
    await database.close()


@pytest.mark.asyncio
async def test_event_cluster_export_contains_candidate_event_and_audit_sheets(tmp_path: Path) -> None:
    database = AsyncDatabase(f"sqlite+aiosqlite:///{tmp_path / 'events.db'}")
    await database.initialize()
    await database.enqueue_job("job-v2", "事件导出", mode="single", total_records=2, pipeline_version="event_cluster_v2")
    first, second = await database.add_records("job-v2", [
        {"source": "S", "source_row": 2, "work_order_id": "A1", "title": "甲公司欠薪", "region": "江海区", "street": "外海街道", "category_level_1": "劳动保障", "category": "拖欠工资"},
        {"source": "S", "source_row": 3, "work_order_id": "A2", "title": "甲公司拖欠工资", "region": "江海区", "street": "外海街道", "category_level_1": "劳动保障", "category": "拖欠工资"},
    ])
    await database.upsert_candidate_pairs("job-v2", [{"record_a_id": first, "record_b_id": second, "recall_reason": "hybrid_vector", "rerank_score": 0.96}])
    await database.replace_candidate_events("job-v2", [{
        "name": "江海区｜外海街道｜甲公司｜拖欠工资",
        "status": "auto_merged",
        "confidence": 0.97,
        "members": [{"record_id": first, "confidence": 0.97}, {"record_id": second, "confidence": 0.96}],
    }])

    output = await export_job(database, "job-v2", tmp_path / "result.xlsx")
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)

    assert workbook.sheetnames == [
        "结果总览", "地区事项统计", "工单明细", "候选事件", "正式事件组", "召回审计", "抽取失败", "模型失败"
    ]
    headers = [cell.value for cell in next(workbook["工单明细"].iter_rows())]
    assert {"候选事件ID", "候选事件名称", "事件状态", "事件置信度", "成员置信度", "合并来源", "事项分类一级"}.issubset(headers)
    await database.close()


@pytest.mark.asyncio
async def test_async_export_escapes_formula_like_untrusted_values(tmp_path: Path) -> None:
    database = AsyncDatabase(f"sqlite+aiosqlite:///{tmp_path / 'formula.db'}")
    await database.initialize()
    await database.enqueue_job("job-formula", "公式转义", mode="single", total_records=1, pipeline_version="event_cluster_v2")
    record_id = (await database.add_records("job-formula", [{
        "source": "S", "source_row": 2, "title": "=HYPERLINK(\"https://example.invalid\")",
        "appeal_text": "+cmd", "raw_json": {"危险列": "@SUM(1,1)"},
    }]))[0]
    await database.replace_candidate_events("job-formula", [{
        "name": "=1+1", "status": "singleton", "members": [{"record_id": record_id}],
    }])

    output = await export_job(database, "job-formula", tmp_path / "formula.xlsx")
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=False)
    rows = list(workbook["工单明细"].iter_rows(values_only=True))
    values = dict(zip(rows[0], rows[1]))

    assert values["候选事件名称"].startswith("'=")
    assert values["诉求标题"].startswith("'=")
    assert values["市民诉求"].startswith("'+")
    assert values["原始字段｜危险列"].startswith("'@")
    await database.close()
