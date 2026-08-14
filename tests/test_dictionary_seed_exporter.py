from pathlib import Path

import pytest
from openpyxl import load_workbook

from complaint_dedup.async_database import AsyncDatabase
from complaint_dedup.corpus_pipeline import CorpusProcessor
from complaint_dedup.corpus_repository import CorpusRepository
from complaint_dedup.dictionary_seed_exporter import export_dictionary_seed
from complaint_dedup.pipeline import InputRecord


@pytest.mark.asyncio
async def test_dictionary_seed_export_contains_review_sheets_and_evidence(tmp_path: Path):
    database = AsyncDatabase(f"sqlite+aiosqlite:///{tmp_path / 'seed.db'}")
    await database.initialize()
    repository = CorpusRepository(database)
    processor = CorpusProcessor(repository)
    records = [
        InputRecord(
            source="B",
            source_row=row,
            work_order_id=f"WO-{row}",
            title="德昌电机门口积水",
            category="道路积水",
            category_level_4="道路积水",
            appeal_text="地址：江海区礼乐街道德昌电机门口。\n事项：道路积水。",
            raw_fields={"工单编号": f"WO-{row}"},
        )
        for row in (2, 3)
    ]
    staged = await processor.stage_records(
        name="历史",
        batch_type="bootstrap_history",
        file_name="history.xlsx",
        file_hash="e" * 64,
        records=records,
    )

    output = await export_dictionary_seed(
        repository, staged.dictionary_version_id, tmp_path / "dictionary_seed_v1.xlsx"
    )
    await database.close()

    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == [
        "标准街道",
        "标准锚点",
        "标准问题",
        "审核队列",
        "抽取统计",
    ]
    anchors = workbook["标准锚点"]
    headers = [cell.value for cell in anchors[1]]
    evidence_column = headers.index("证据数") + 1
    assert anchors.cell(2, evidence_column).value == 2
