import asyncio
from io import BytesIO
from pathlib import Path
import time

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl import load_workbook

from complaint_dedup.async_database import AsyncDatabase
from complaint_dedup.config import Settings
from complaint_dedup.corpus_web import create_corpus_app


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "工单编号",
            "受理时间",
            "诉求标题",
            "事发地点",
            "事项分类四级",
            "事项分类",
            "所属部门",
            "市民诉求",
            "联系电话",
        ]
    )
    sheet.append(
        [
            "WO-1",
            "2026-08-12 08:00:00",
            "德昌电机门口积水",
            "江海区礼乐街道德昌电机门口",
            "道路积水",
            "道路积水",
            "礼乐街道办事处",
            "地址：江海区礼乐街道德昌电机门口。\n事项：道路积水。",
            "13800138000",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_client(tmp_path: Path) -> TestClient:
    settings = Settings(
        database_mode="sqlite",
        database_path=tmp_path / "app.db",
        max_total_rows=200_000,
        _env_file=None,
    )
    database = AsyncDatabase(f"sqlite+aiosqlite:///{settings.database_path}")
    return TestClient(create_corpus_app(settings, database=database))


def wait_for_text(client: TestClient, url: str, expected: str, timeout: float = 5) -> str:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        text = client.get(url).text
        if expected in text:
            return text
        time.sleep(0.05)
    raise AssertionError(f"等待页面出现文本超时: {expected}")


def test_department_filter_submits_immediately(tmp_path: Path):
    with make_client(tmp_path) as client:
        response = client.get("/events")

    assert response.status_code == 200

    assert 'streetFilter.value = ""' in response.text

    assert 'select name="processing_department" id="event-department-filter"' in response.text
    assert 'input type="date" name="completed_from"' in response.text
    assert 'input type="date" name="completed_to"' in response.text
    assert 'checkbox" name="missing_completed" value="1"' in response.text
    assert 'departmentFilter.addEventListener("change"' in response.text
    assert 'select name="sort" id="event-sort-filter"' in response.text
    assert "工单数从多到少" in response.text
    assert "工单数从少到多" in response.text
    assert 'sortFilter.addEventListener("change"' in response.text
    assert 'input type="checkbox" name="has_daily" value="1"' in response.text
    assert 'input type="checkbox" name="hide_singletons" value="1"' in response.text
    assert "单条工单事件" in response.text


def test_mobile_forms_and_batch_preview_have_bounded_fields_and_card_labels():
    css = Path("static/app.css").read_text(encoding="utf-8")
    batch_template = Path("templates/corpus_batch.html").read_text(encoding="utf-8")

    assert "label > input, label > select" in css
    assert "min-width: 0" in css
    assert 'class="batch-preview-table"' in batch_template
    assert 'data-label="工单编号"' in batch_template


def test_upload_is_queued_without_parsing_in_request(
    tmp_path: Path, monkeypatch
):
    settings = Settings(
        database_mode="sqlite",
        database_path=tmp_path / "queued.db",
        max_total_rows=200_000,
        _env_file=None,
    )
    database = AsyncDatabase(f"sqlite+aiosqlite:///{settings.database_path}")

    async def fail_if_called(*args, **kwargs):
        raise AssertionError("上传请求不应解析 Excel")

    app = create_corpus_app(settings, database=database, embedded_worker=False)
    monkeypatch.setattr(app.state.processor, "stage_records", fail_if_called)
    with TestClient(app) as client:
        response = client.post(
            "/batches",
            data={"mode": "bootstrap_history", "name": "异步历史"},
            files={
                "file_history": (
                    "history.xlsx",
                    workbook_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        page = client.get(response.headers["location"])
        assert "等待后台处理" in page.text


def test_home_shows_corpus_batch_modes(tmp_path: Path):
    with make_client(tmp_path) as client:
        response = client.get("/")
    assert response.status_code == 200
    assert "历史库冷启动" in response.text
    assert "首次 A/B 联合比对" in response.text
    assert "每日新增" in response.text
    assert "导入批次" in response.text
    assert "历史批次" not in response.text
    assert "补录或更正" not in response.text
    assert "词典状态" not in response.text
    assert "词典审核" not in response.text


def test_batch_history_is_paginated_by_ten(tmp_path: Path):
    settings = Settings(
        database_mode="sqlite",
        database_path=tmp_path / "paged.db",
        max_total_rows=200_000,
        _env_file=None,
    )
    database = AsyncDatabase(f"sqlite+aiosqlite:///{settings.database_path}")
    app = create_corpus_app(settings, database=database, embedded_worker=False)
    with TestClient(app) as client:
        for index in range(12):
            asyncio.run(
                app.state.repository.create_batch(
                    f"分页批次-{index:02d}", "daily_increment"
                )
            )

        first_page = client.get("/batches")
        second_page = client.get("/batches?page=2")
        home = client.get("/")

    assert first_page.status_code == 200
    assert first_page.text.count("分页批次-") == 10
    assert "第 1 页，共 2 页" in first_page.text
    assert second_page.text.count("分页批次-") == 2
    assert "第 2 页，共 2 页" in second_page.text
    assert "<strong>12</strong>" in home.text


def test_correction_mode_is_rejected(tmp_path: Path):
    with make_client(tmp_path) as client:
        response = client.post(
            "/batches",
            data={"mode": "correction", "name": "不应创建"},
            follow_redirects=False,
        )
    assert response.status_code == 400
    assert "批次类型无效" in response.text


def test_dictionary_routes_are_hidden_from_users(tmp_path: Path):
    with make_client(tmp_path) as client:
        response = client.get("/batches/unknown/dictionary")
    assert response.status_code == 404


def test_bootstrap_upload_auto_commits_and_event_page(tmp_path: Path):
    with make_client(tmp_path) as client:
        response = client.post(
            "/batches",
            data={"mode": "bootstrap_history", "name": "历史测试"},
            files={
                "file_history": (
                    "history.xlsx",
                    workbook_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        batch_url = response.headers["location"]
        batch_text = wait_for_text(client, batch_url, "批次已完成")
        assert "历史库冷启动" in batch_text
        assert "bootstrap_history" not in batch_text
        assert "挂入已有事件" in batch_text
        assert "新增事件" in batch_text
        assert "<span>状态</span>" not in batch_text
        assert "<span>阶段</span>" not in batch_text
        assert client.get(f"{batch_url}/dictionary?dimension=anchor").status_code == 404
        events = client.get("/events")
        assert "德昌电机门口" in events.text
        assert 'data-label="事件名称"' in events.text
        assert '<option value="江海区"' in events.text
        assert '<option value="礼乐街道"' in events.text

        assert "/events/1" in events.text
        assert "德昌电机门口" in client.get(
            "/events?region=江海区&street=礼乐街道&event_name=德昌电机门口"
        ).text
        assert "德昌电机门口" in client.get(
            "/events?sort=member_count_desc"
        ).text
        daily_filtered = client.get(
            "/events?has_daily=1&hide_singletons=1&sort=member_count_desc"
        )
        assert daily_filtered.status_code == 200
        assert 'name="has_daily" value="1" checked' in daily_filtered.text
        assert 'name="hide_singletons" value="1" checked' in daily_filtered.text
        pagination = client.get("/events?page=1").text
        assert "has_daily=&" not in pagination
        assert "hide_singletons=&" not in pagination
        stale_browser_url = client.get(
            "/events?region=&street=&event_name=&sort=updated_desc"
            "&has_daily=&hide_singletons=1&page=2"
        )
        assert stale_browser_url.status_code == 200
        assert 'name="hide_singletons" value="1" checked' in stale_browser_url.text
        assert "德昌电机门口" in client.get("/events?event_name=德昌电机").text
        assert "德昌电机门口" not in client.get("/events?event_name=不存在").text
        detail = client.get("/events/1")
        assert detail.status_code == 200
        assert "受理时间" in detail.text
        assert "所属部门" in detail.text
        assert "数据来源" in detail.text
        assert "成员置信度" not in detail.text
        assert 'data-label="受理时间"' in detail.text
        assert 'data-label="操作"' in detail.text
        assert "/events/options?q=" in detail.text
        options = client.get("/events/options?q=德昌")
        assert options.status_code == 200
        assert options.json() == [
            {"id": 1, "event_name": "江海区｜礼乐街道｜德昌电机门口｜道路积水"}
        ]

        renamed = client.post(
            "/events/1/name",
            data={"name": "礼乐街道｜德昌电机门口｜道路积水"},
            follow_redirects=False,
        )
        assert renamed.status_code == 303
        assert "礼乐街道｜德昌电机门口｜道路积水" in client.get("/events/1").text
        exported = client.get("/exports/corpus")
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.content))
        assert workbook.sheetnames == ["重复项", "孤立工单"]
