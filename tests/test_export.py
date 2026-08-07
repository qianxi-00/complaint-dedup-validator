from pathlib import Path

from openpyxl import load_workbook

from complaint_dedup.database import connect_database, initialize_database
from complaint_dedup.exporter import export_job


def test_export_contains_required_worksheets(tmp_path: Path) -> None:
    database = tmp_path / "app.db"
    output = tmp_path / "result.xlsx"
    initialize_database(database)
    with connect_database(database) as connection:
        connection.execute(
            "INSERT INTO jobs (id, status, stage, total_records) VALUES ('job', 'review_ready', 'review', 2)"
        )
        connection.execute(
            """
            INSERT INTO records (job_id, source, source_row, work_order_id, title, raw_json, extraction_status)
            VALUES ('job', 'A', 2, 'A1', '甲公司欠薪', '{}', 'succeeded')
            """
        )
        connection.execute(
            """
            INSERT INTO records (job_id, source, source_row, work_order_id, title, raw_json, extraction_status)
            VALUES ('job', 'B', 2, 'B1', '甲公司拖欠工资', '{}', 'succeeded')
            """
        )
        connection.execute(
            """
            INSERT INTO candidate_pairs (
                job_id, record_a_id, record_b_id, judgement_status, llm_decision, confidence
            ) VALUES ('job', 1, 2, 'succeeded', 'duplicate', 0.95)
            """
        )
        connection.execute(
            "INSERT INTO reviews (candidate_pair_id, decision) VALUES (1, 'duplicate')"
        )
        group = connection.execute(
            "INSERT INTO event_groups (job_id, name) VALUES ('job', '甲公司｜欠薪')"
        )
        connection.executemany(
            "INSERT INTO event_members (event_group_id, record_id) VALUES (?, ?)",
            [(group.lastrowid, 1), (group.lastrowid, 2)],
        )
        connection.execute(
            """
            INSERT INTO llm_batches (
                job_id, batch_type, batch_index, status, attempts,
                request_json, response_json, error_message
            ) VALUES ('job', 'judgement', 0, 'failed', 3, '[1]', 'raw output', 'invalid json')
            """
        )

    export_job(database, "job", output)

    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == ["结果总览", "候选对", "事件组", "抽取失败", "模型失败"]
    failure_rows = list(workbook["模型失败"].iter_rows(values_only=True))
    assert "response_json" in failure_rows[0]
    assert "raw output" in failure_rows[1]


def test_export_escapes_formula_like_values(tmp_path: Path) -> None:
    database = tmp_path / "app.db"
    output = tmp_path / "formula.xlsx"
    initialize_database(database)
    with connect_database(database) as connection:
        connection.execute(
            "INSERT INTO jobs (id, status, stage, total_records) VALUES ('job', 'review_ready', 'review', 1)"
        )
        connection.execute(
            "INSERT INTO records (job_id, source, source_row, title, appeal_text, raw_json) VALUES ('job', 'A', 2, '=1+1', '@cmd', '{}')"
        )
        connection.execute(
            "INSERT INTO records (job_id, source, source_row, title, appeal_text, raw_json) VALUES ('job', 'B', 2, '普通标题', '普通内容', '{}')"
        )
        connection.execute(
            "INSERT INTO candidate_pairs (job_id, record_a_id, record_b_id, candidate_key) VALUES ('job', 1, 2, 'test')"
        )
    export_job(database, "job", output)
    workbook = load_workbook(output, data_only=False, read_only=True)
    values = list(workbook["候选对"].iter_rows(values_only=True))
    flattened = [value for row in values for value in row if value is not None]
    assert "'=1+1" in flattened
    assert "'@cmd" in flattened
