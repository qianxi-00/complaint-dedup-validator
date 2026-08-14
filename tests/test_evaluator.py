import pytest
import pandas as pd
from openpyxl import load_workbook

from complaint_dedup.evaluator import evaluate_explicit_pairs
from complaint_dedup.evaluator import export_evaluation, load_explicit_pairs
from complaint_dedup.llm_models import JudgementBatchResponse


class FakeClient:
    def __init__(self) -> None:
        self.calls = []

    async def chat_json(self, messages, response_model):
        self.calls.append(messages)
        pair_id = "1|2" if len(self.calls) == 1 else "3|4"
        return JudgementBatchResponse.model_validate(
            {
                "pairs": [
                    {
                        "pair_id": pair_id,
                        "decision": "review",
                        "confidence": 0.5,
                        "subject_relation": "unknown",
                        "address_relation": "unknown",
                        "issue_relation": "unknown",
                        "new_independent_issue": False,
                        "hard_conflicts": [],
                        "evidence_a": [],
                        "evidence_b": [],
                        "reason": "信息不足",
                        "matrix": {},
                    }
                ]
            }
        )


@pytest.mark.asyncio
async def test_explicit_evaluator_preserves_requested_pairs_and_batches() -> None:
    client = FakeClient()
    pairs = [
        {"pair_id": "1|2", "a": {"title": "甲"}, "b": {"title": "乙"}},
        {"pair_id": "3|4", "a": {"title": "丙"}, "b": {"title": "丁"}},
    ]

    result = await evaluate_explicit_pairs(client, pairs, batch_size=1)

    assert [item.pair_id for item in result] == ["1|2", "3|4"]
    assert len(client.calls) == 2


@pytest.mark.asyncio
async def test_explicit_evaluator_rejects_missing_pair_result() -> None:
    client = FakeClient()
    pairs = [{"pair_id": "missing", "a": {}, "b": {}}]

    with pytest.raises(ValueError, match="missing"):
        await evaluate_explicit_pairs(client, pairs, batch_size=10)


def test_load_explicit_pairs_from_excel_preserves_reference_label(tmp_path) -> None:
    path = tmp_path / "pairs.xlsx"
    pd.DataFrame(
        [
            {
                "pair_id": "P001",
                "A工单编号": "A1",
                "A标题": "甲公司车辆故障",
                "A内容": "金瓯路188号，要求维修",
                "B工单编号": "B1",
                "B标题": "甲公司车辆抖动",
                "B内容": "金瓯路188号，要求维修",
                "参考标签": "duplicate",
            }
        ]
    ).to_excel(path, index=False)

    rows = load_explicit_pairs(path)

    assert rows[0]["pair_id"] == "P001"
    assert rows[0]["a"]["work_order_id"] == "A1"
    assert rows[0]["reference_label"] == "duplicate"


def test_export_evaluation_contains_metrics_and_pair_details(tmp_path) -> None:
    output = tmp_path / "result.xlsx"
    source_pairs = [
        {
            "pair_id": "P001",
            "a": {"title": "甲"},
            "b": {"title": "乙"},
            "reference_label": "review",
        }
    ]
    judged = JudgementBatchResponse.model_validate(
        {
            "pairs": [
                {
                    "pair_id": "P001",
                    "decision": "review",
                    "confidence": 0.5,
                    "subject_relation": "unknown",
                    "address_relation": "unknown",
                    "issue_relation": "unknown",
                    "new_independent_issue": False,
                    "hard_conflicts": [],
                    "evidence_a": [],
                    "evidence_b": [],
                    "reason": "信息不足",
                    "matrix": {"same_object": None},
                }
            ]
        }
    ).pairs

    export_evaluation(source_pairs, judged, output)

    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == ["评测总览", "逐对结果"]
    assert workbook["逐对结果"].max_row == 2
